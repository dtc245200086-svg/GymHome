import os
import sys

try:
    from docx import Document
except ModuleNotFoundError:
    print("Module 'python-docx' chưa được cài. Chạy: pip install python-docx")
    sys.exit(1)

path = sys.argv[1] if len(sys.argv) > 1 else r'Báo Cáo SRS-Hệ Thống Quản Lý Phòng Gym.docx'

if not os.path.exists(path):
    print(f"File không tồn tại: {path}")
    print("Sử dụng: python read_docx.py <path/to/file.docx>")
    sys.exit(1)

print(f"Đang đọc file: {path}")

doc = Document(path)

print('-' * 60)
print('--- Nội dung đoạn văn ---')
print('-' * 60)
for para in doc.paragraphs:
    text = para.text.strip()
    if text:
        print(text)

print('-' * 60)
print('--- Nội dung bảng ---')
print('-' * 60)
for table_idx, table in enumerate(doc.tables, start=1):
    print(f"Bảng {table_idx}:")
    for row in table.rows:
        row_text = ' | '.join(cell.text.strip() for cell in row.cells)
        print(row_text)
    print('-' * 40)

print('Hoàn tất đọc file.')